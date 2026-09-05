from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

path = r"C:\StudyMind AI\output\documents\StudyMind_AI_Mini_Project_Synopsis.docx"
doc = Document(path)
sections = {
    "Objectives of the Project": 0,
    "User Flow": 0,
    "Expected Outcomes": 0,
    "References": 0,
}
active = None
for p in doc.paragraphs:
    text = p.text.strip()
    if text in sections:
        active = text
        continue
    if active and p.style.name == "List Number":
        sections[active] += 1
        p.style = doc.styles["Normal"]
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(4 if active != "References" else 7)
        p.paragraph_format.line_spacing = 1.1
        p.clear()
        run = p.add_run(f"{sections[active]}. {text}")
        run.font.name = "Aptos"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        run.font.size = Pt(10.5)
    elif active and text and p.style.name != "List Number":
        active = None
doc.save(path)
